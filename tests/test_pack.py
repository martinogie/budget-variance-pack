import unittest
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from budget_pack import analysis as A
from budget_pack.data import CLOSE_MONTH, build

XLSX = Path(__file__).resolve().parent.parent / "reports" / "FY2026_Budget_vs_Actual_Variance_Pack.xlsx"


class Data:
    gl, drv, hc = build()
    pnl = A.pnl(gl)


class TestAnalysis(unittest.TestCase):
    def test_bridge_adds_up_per_product(self):
        g = A.price_volume_mix(Data.drv, range(1, CLOSE_MONTH + 1))
        np.testing.assert_allclose(g["volume"] + g["mix"] + g["rate"], g["total"], atol=1e-4)

    def test_bridge_total_equals_product_revenue_variance(self):
        g = A.price_volume_mix(Data.drv, range(1, CLOSE_MONTH + 1))
        self.assertAlmostEqual(g["total"].sum(), Data.pnl.loc[list(A.PRODUCTS), "ytd_var"].sum(), delta=1)

    def test_headcount_effects_add_up(self):
        h = A.headcount_bridge(Data.gl, Data.hc)
        np.testing.assert_allclose(h["volume_effect"] + h["rate_effect"], h["ytd_var"], atol=1e-6)

    def test_ebitda_is_revenue_less_costs(self):
        p = Data.pnl
        for c in ("m_act", "ytd_bud", "fy_out"):
            self.assertAlmostEqual(p.loc["EBITDA", c], p.loc["Total revenue", c] - p.loc["Total cost of revenue", c]
                                   - p.loc["Total operating expenses", c], delta=0.01)

    def test_favourable_sign_convention(self):
        p = Data.pnl
        self.assertGreater(p.loc["Marketing programs", "ytd_var"], 0)          # under budget = favourable
        self.assertLess(p.loc["Fraud & chargeback losses", "ytd_var"], 0)      # over budget = unfavourable


@unittest.skipUnless(XLSX.exists(), "run `python -m budget_pack.run --recalc` first")
class TestWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(XLSX, data_only=True)
        ws = cls.wb["P&L Summary"]
        cls.rows = {r[0]: r for r in ws.iter_rows(min_row=6, values_only=True) if r[0]}
        if cls.rows["Total revenue"][2] is None:
            raise unittest.SkipTest("workbook has no cached values; recalculate it first")

    def test_no_formula_errors(self):
        for ws in self.wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    self.assertFalse(isinstance(v, str) and v.startswith("#") and v.endswith(("!", "?", "A")), (ws.title, v))

    def test_pnl_matches_pandas(self):
        cols = {2: "m_act", 3: "m_bud", 4: "m_var", 6: "ytd_act", 7: "ytd_bud", 8: "ytd_var", 10: "fy_bud", 11: "fy_out", 12: "fy_var"}
        for line, row in self.rows.items():
            if line not in Data.pnl.index:
                continue
            for i, col in cols.items():
                self.assertAlmostEqual(row[i], Data.pnl.loc[line, col], delta=1.0, msg=f"{line} {col}")

    def test_bridge_matches_pandas(self):
        g = A.price_volume_mix(Data.drv, range(1, CLOSE_MONTH + 1))
        ws = self.wb["Revenue Bridge"]
        tot = [r for r in ws.iter_rows(values_only=True) if r[0] == "Total products"][-1]   # YTD block
        self.assertAlmostEqual(tot[8], g["volume"].sum(), delta=1)
        self.assertAlmostEqual(tot[9], g["mix"].sum(), delta=1)
        self.assertAlmostEqual(tot[10], g["rate"].sum(), delta=1)

    def test_headcount_matches_pandas(self):
        h = A.headcount_bridge(Data.gl, Data.hc).set_index("department")
        ws = self.wb["Headcount"]
        for r in ws.iter_rows(min_row=6, values_only=True):
            if r[0] in h.index:
                self.assertAlmostEqual(r[10], h.loc[r[0], "volume_effect"], delta=1)
                self.assertAlmostEqual(r[11], h.loc[r[0], "rate_effect"], delta=1)


if __name__ == "__main__":
    unittest.main()
