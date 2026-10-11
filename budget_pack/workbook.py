"""Write the variance pack as a formula-driven Excel workbook.

Every number in the report tabs is a formula over the data tabs (GL_Data, Drivers, HC_Data), so
changing the reporting month on the Control tab, or a forecast input, updates the whole pack.
"""
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .analysis import FORECAST, REVENUE, THRESHOLD_PCT, THRESHOLD_USD
from .data import CLOSE_MONTH, COGS, DEPTS, OPEX, PRODUCTS, YEAR

FONT = "Arial"
NAVY = "1F3864"
F_TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_BODY = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_LINK_BOLD = Font(name=FONT, size=10, color="008000", bold=True)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_TOTAL = PatternFill("solid", fgColor="F2F2F2")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
TOP = Border(top=Side(style="thin", color="000000"))
USD = '$#,##0;[Red]($#,##0);"-"'
USD_PLAIN = '$#,##0;($#,##0);"-"'
PCT = '0.0%;[Red](0.0%);"-"'
RATE = '0.000%'
NUM = '#,##0;(#,##0);"-"'
FACTOR = '0.000"x"'


class Pack:
    def __init__(self, gl, drv, hc, commentary):
        self.gl_df, self.drv_df, self.hc_df, self.commentary = gl, drv, hc, commentary
        self.wb = Workbook()
        self.n_gl = len(gl) + 1
        self.n_drv = len(drv) + 1
        self.n_hc = len(hc) + 1
        self.fc_rows = {}

    # ---------- range helpers ----------
    def gl(self, col):
        c = {"scenario": "A", "month": "B", "month_num": "C", "department": "D", "section": "E",
             "line": "F", "type": "G", "amount": "H"}[col]
        return f"GL_Data!${c}$2:${c}${self.n_gl}"

    def drv(self, col):
        c = {"month_num": "B", "product": "C", "budget_tpv": "D", "budget_rate": "E",
             "actual_tpv": "F", "actual_rate": "G"}[col]
        return f"Drivers!${c}$2:${c}${self.n_drv}"

    def hcr(self, col):
        c = {"month_num": "B", "department": "C", "budget_hc": "D", "actual_hc": "E", "cph": "F"}[col]
        return f"HC_Data!${c}$2:${c}${self.n_hc}"

    def gl_sum(self, scenario, period, line=None, dept=None):
        """period: 'm' (reporting month), 'ytd', 'fy', or a (lo, hi) tuple of fixed months."""
        crit = [f'{self.gl("scenario")},"{scenario}"']
        if line:
            crit.append(f'{self.gl("line")},{line}')
        if dept:
            crit.append(f'{self.gl("department")},{dept}')
        crit += self._period(self.gl("month_num"), period)
        return f'SUMIFS({self.gl("amount")},{",".join(crit)})'

    @staticmethod
    def _period(rng, period):
        if period == "m":
            return [f"{rng},Control!$C$4"]
        if period == "ytd":
            return [f'{rng},"<="&Control!$C$4']
        if period == "fy":
            return []
        lo, hi = period
        return [f'{rng},">={lo}"', f'{rng},"<={hi}"']

    # ---------- styling helpers ----------
    @staticmethod
    def head(ws, row, labels, start=1):
        for i, text in enumerate(labels):
            c = ws.cell(row=row, column=start + i, value=text)
            c.font, c.fill = F_HEAD, FILL_HEAD
            c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    @staticmethod
    def widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    @staticmethod
    def title(ws, text, sub):
        ws["A1"] = text
        ws["A1"].font = F_TITLE
        ws["A2"] = sub
        ws["A2"].font = F_SUB
        ws.sheet_view.showGridLines = False

    @staticmethod
    def setup_print(ws):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    def flag_rules(self, ws, rng):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Unfav"'],
                                                      fill=PatternFill("solid", fgColor="F8CBAD"), font=Font(name=FONT, color="9C0006", bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fav"'],
                                                      fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=FONT, color="006100", bold=True)))

    @staticmethod
    def var(type_cell, act, bud):
        return f'=IF({type_cell}="Revenue",{act}-{bud},{bud}-{act})'

    @staticmethod
    def pct(var, bud):
        return f'=IF({bud}=0,"",{var}/ABS({bud}))'

    @staticmethod
    def flag(var, pct):
        return (f'=IFERROR(IF(AND(ABS({var})>=Control!$C$5,ABS({pct})>=Control!$C$6),'
                f'IF({var}>0,"Fav","Unfav"),""),"")')

    # ---------- tabs ----------
    def control(self):
        ws = self.wb.active
        ws.title = "Control"
        self.title(ws, f"FY{YEAR} Budget vs Actual Variance Pack",
                   "Card-payments company (synthetic data). All report tabs are formulas over the data tabs.")
        rows = [
            ("Reporting month (1-9)", CLOSE_MONTH, "Change this to re-run the pack for any closed month.", True),
            ("Flag variances of at least ($)", THRESHOLD_USD, "Both thresholds must be met to flag a line.", False),
            ("...and at least (%)", THRESHOLD_PCT, None, False),
        ]
        for i, (label, value, note, key) in enumerate(rows, 4):
            ws.cell(row=i, column=2, value=label).font = F_BODY
            c = ws.cell(row=i, column=3, value=value)
            c.font = F_INPUT
            if key:
                c.fill = FILL_KEY
            if note:
                ws.cell(row=i, column=5, value=note).font = F_SUB
        ws["C5"].number_format = USD_PLAIN
        ws["C6"].number_format = "0%"
        ws["D4"] = f'=TEXT(DATE({YEAR},C4,1),"mmmm yyyy")'
        ws["D4"].font = F_BOLD
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(CLOSE_MONTH),
                            showErrorMessage=True, error=f"Actuals are loaded for months 1-{CLOSE_MONTH}.")
        ws.add_data_validation(dv)
        dv.add("C4")
        ws["C4"].comment = Comment(f"Input. Actuals are loaded through month {CLOSE_MONTH} (September).", "Pack")

        ws["B9"] = "How to read this pack"
        ws["B9"].font = F_BOLD
        guide = [
            ("P&L Summary", "Month, year-to-date and full-year outlook vs budget, with favourable/unfavourable flags."),
            ("Dept Opex", "Operating expenses by department and account for the month and YTD."),
            ("Revenue Bridge", "Product revenue variance split into volume, mix and rate effects, with a waterfall chart."),
            ("Headcount", "Salary variance split into headcount (volume) and cost-per-head (rate)."),
            ("Forecast", "Full-year outlook: YTD actual + Q4 forecast. Method and adjustments are inputs (blue)."),
            ("Commentary", "Written explanation of the September close."),
            ("GL_Data / Drivers / HC_Data", "Source data: GL by month, department and line; payment volume and take rates; headcount."),
        ]
        for i, (tab, text) in enumerate(guide, 10):
            ws.cell(row=i, column=2, value=tab).font = F_BOLD
            ws.cell(row=i, column=3, value=text).font = F_BODY
        ws["B18"] = "Conventions"
        ws["B18"].font = F_BOLD
        legend = [
            (F_INPUT, None, "Blue text: inputs you can change"),
            (F_INPUT, FILL_KEY, "Yellow fill: key assumption"),
            (F_BODY, None, "Black text: formulas"),
            (F_LINK, None, "Green text: links to another tab"),
            (F_BODY, None, "Fav/(Unfav): positive = favourable (revenue above budget, or cost below budget)"),
        ]
        for i, (font, fill, text) in enumerate(legend, 19):
            c = ws.cell(row=i, column=2, value=text)
            c.font = font
            if fill:
                c.fill = fill
        self.widths(ws, [2, 30, 92, 16, 50])
        self.setup_print(ws)

    def data_tabs(self):
        specs = [
            ("GL_Data", self.gl_df, ["scenario", "month", "month_num", "department", "section", "line", "type", "amount"],
             ["Scenario", "Month", "Month #", "Department", "Section", "Line", "Type", "Amount ($)"], [10, 10, 9, 20, 10, 30, 10, 14]),
            ("Drivers", self.drv_df, ["month", "month_num", "product", "budget_tpv", "budget_rate", "actual_tpv", "actual_rate"],
             ["Month", "Month #", "Product", "Budget TPV ($)", "Budget take rate", "Actual TPV ($)", "Actual take rate"], [10, 9, 30, 16, 14, 16, 14]),
            ("HC_Data", self.hc_df, ["month", "month_num", "department", "budget_hc", "actual_hc", "budget_cost_per_head"],
             ["Month", "Month #", "Department", "Budget HC", "Actual HC", "Budget cost per head ($/month)"], [10, 9, 20, 11, 11, 16]),
        ]
        for name, df, cols, labels, widths in specs:
            ws = self.wb.create_sheet(name)
            self.head(ws, 1, labels)
            for r, rec in enumerate(df[cols].itertuples(index=False), 2):
                for c, v in enumerate(rec, 1):
                    cell = ws.cell(row=r, column=c, value=None if v != v else v)
                    cell.font = F_BODY
            fmt = {"amount": USD_PLAIN, "budget_tpv": USD_PLAIN, "actual_tpv": USD_PLAIN, "budget_rate": RATE,
                   "actual_rate": RATE, "budget_cost_per_head": USD_PLAIN}
            for c, col in enumerate(cols, 1):
                if col in fmt:
                    for r in range(2, len(df) + 2):
                        ws.cell(row=r, column=c).number_format = fmt[col]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df) + 1}"
            self.widths(ws, widths)
            ws.sheet_properties.tabColor = "808080"

    def pnl_structure(self):
        s = [("section", "Revenue")]
        s += [("line", l, "Revenue") for l in REVENUE]
        s += [("total", "Total revenue", "Revenue", REVENUE)]
        s += [("section", "Cost of revenue")]
        s += [("line", l, "Cost") for l in COGS]
        s += [("total", "Total cost of revenue", "Cost", COGS)]
        s += [("calc", "Gross profit", "Revenue", ("Total revenue", "Total cost of revenue"))]
        s += [("margin", "Gross margin %", ("Gross profit", "Total revenue"))]
        s += [("section", "Operating expenses")]
        s += [("line", l, "Cost") for l in OPEX]
        s += [("total", "Total operating expenses", "Cost", OPEX)]
        s += [("calc3", "EBITDA", "Revenue", ("Gross profit", "Total operating expenses"))]
        s += [("margin", "EBITDA margin %", ("EBITDA", "Total revenue"))]
        return s

    def forecast_tab(self):
        ws = self.wb.create_sheet("Forecast")
        self.title(ws, "Full-year outlook (as at the September close)",
                   "Outlook = Jan-Sep actual + Q4 forecast. Q4 forecast = Q4 budget x run-rate factor (Q3 actual / Q3 budget) "
                   "or Q4 budget, plus a $ adjustment. Method and adjustment are inputs.")
        cols = ["Line", "Type", "Q4 method", "Q4 adjustment ($)", "Jan-Sep actual", "Q3 actual", "Q3 budget",
                "Run-rate factor", "Q4 budget", "Q4 forecast", "FY outlook", "FY budget", "Fav/(Unfav)", "Why"]
        self.head(ws, 4, cols)
        dv = DataValidation(type="list", formula1='"Run-rate,Budget"', showErrorMessage=True)
        ws.add_data_validation(dv)
        notes = {
            "Fraud & chargeback losses": "August card-testing attack treated as one-off; $150K for extra monitoring",
            "Marketing programs": "Q3 campaign ($1.08M) deferred into Q4",
            "Professional fees": "Money-transmitter licensing continues into Q4",
            "Platform subscriptions": "No change to plan; renewals as budgeted",
            "Card processing - Enterprise": "Enterprise wins ramp; negotiated rate carries into Q4",
            "Instant payouts (RTP)": "Adoption ahead of plan",
            "Contractors": "Engineering backfill continues until roles are filled",
        }
        r = 5
        rows = {}
        for item in self.pnl_structure():
            kind, name = item[0], item[1]
            if kind == "section":
                ws.cell(row=r, column=1, value=name).font = F_BOLD
                for c in range(1, 15):
                    ws.cell(row=r, column=c).fill = FILL_SECTION
                r += 1
                continue
            if kind == "margin":
                continue
            t = item[2]
            ws.cell(row=r, column=1, value=name).font = F_BOLD if kind != "line" else F_BODY
            ws.cell(row=r, column=2, value=t).font = F_BODY
            if kind == "line":
                method, adj = FORECAST[name]
                ws.cell(row=r, column=3, value=method).font = F_INPUT
                dv.add(f"C{r}")
                ws.cell(row=r, column=4, value=adj).font = F_INPUT
                q = f"$A{r}"
                ws.cell(row=r, column=5, value=f"={self.gl_sum('Actual', (1, CLOSE_MONTH), q)}")
                ws.cell(row=r, column=6, value=f"={self.gl_sum('Actual', (7, 9), q)}")
                ws.cell(row=r, column=7, value=f"={self.gl_sum('Budget', (7, 9), q)}")
                ws.cell(row=r, column=8, value=f'=IF(AND(C{r}="Run-rate",G{r}<>0),F{r}/G{r},1)')
                ws.cell(row=r, column=9, value=f"={self.gl_sum('Budget', (10, 12), q)}")
                ws.cell(row=r, column=10, value=f"=I{r}*H{r}+D{r}")
                ws.cell(row=r, column=11, value=f"=E{r}+J{r}")
                ws.cell(row=r, column=12, value=f"={self.gl_sum('Budget', 'fy', q)}")
                ws.cell(row=r, column=14, value=notes.get(name, "")).font = F_SUB
            else:
                if kind == "total":
                    first, last = rows[item[3][0]], rows[item[3][-1]]
                    f = lambda col: f"=SUM({col}{first}:{col}{last})"
                else:
                    a, b = rows[item[3][0]], rows[item[3][1]]
                    f = lambda col: f"={col}{a}-{col}{b}"
                for c in (5, 6, 7, 9, 10, 11, 12):
                    col = get_column_letter(c)
                    ws.cell(row=r, column=c, value=f(col))
                for c in range(1, 15):
                    ws.cell(row=r, column=c).fill = FILL_TOTAL
                    ws.cell(row=r, column=c).border = TOP
            ws.cell(row=r, column=13, value=self.var(f"$B{r}", f"K{r}", f"L{r}"))
            for c in range(4, 14):
                cell = ws.cell(row=r, column=c)
                cell.number_format = FACTOR if c == 8 else USD
                if c > 4:
                    cell.font = F_BOLD if kind != "line" else F_BODY
            rows[name] = r
            r += 1
        self.fc_rows = rows
        ws.freeze_panes = "C5"
        self.widths(ws, [30, 9, 11, 14, 14, 13, 13, 10, 13, 13, 14, 14, 13, 58])
        self.setup_print(ws)

    def pnl_tab(self):
        ws = self.wb.create_sheet("P&L Summary", 1)
        self.title(ws, "P&L Summary: actual vs budget", None)
        ws["A2"] = '="Reporting month: "&Control!D4&"   |   $ USD   |   Fav/(Unfav): positive is favourable"'
        ws["A2"].font = F_SUB
        for rng, text in (("C4:F4", "Month"), ("G4:J4", "Year to date"), ("K4:M4", "Full year")):
            ws.merge_cells(rng)
            c = ws[rng.split(":")[0]]
            c.value, c.font, c.fill = text, F_HEAD, FILL_HEAD
            c.alignment = Alignment(horizontal="center")
        self.head(ws, 5, ["Line", "Type", "Actual", "Budget", "Fav/(Unfav)", "%", "Actual", "Budget", "Fav/(Unfav)", "%",
                          "Budget", "Outlook", "Fav/(Unfav)", "Flag (month)"])
        r = 6
        rows = {}
        for item in self.pnl_structure():
            kind, name = item[0], item[1]
            if kind == "section":
                ws.cell(row=r, column=1, value=name).font = F_BOLD
                for c in range(1, 15):
                    ws.cell(row=r, column=c).fill = FILL_SECTION
                r += 1
                continue
            if kind == "margin":
                num, den = rows[item[2][0]], rows[item[2][1]]
                ws.cell(row=r, column=1, value=name).font = Font(name=FONT, size=10, italic=True)
                for c in (3, 4, 7, 8, 11, 12):
                    col = get_column_letter(c)
                    ws.cell(row=r, column=c, value=f"=IFERROR({col}{num}/{col}{den},0)").number_format = "0.0%"
                for c, (a, b) in ((5, ("C", "D")), (9, ("G", "H")), (13, ("L", "K"))):
                    cell = ws.cell(row=r, column=c, value=f"={a}{r}-{b}{r}")
                    cell.number_format = '+0.0%;[Red]-0.0%;"-"'
                for c in range(1, 15):
                    ws.cell(row=r, column=c).font = Font(name=FONT, size=10, italic=True)
                r += 1
                continue
            t = item[2]
            ws.cell(row=r, column=1, value=name).font = F_BOLD if kind != "line" else F_BODY
            ws.cell(row=r, column=2, value=t).font = F_BODY
            if kind == "line":
                q = f"$A{r}"
                ws.cell(row=r, column=3, value=f"={self.gl_sum('Actual', 'm', q)}")
                ws.cell(row=r, column=4, value=f"={self.gl_sum('Budget', 'm', q)}")
                ws.cell(row=r, column=7, value=f"={self.gl_sum('Actual', 'ytd', q)}")
                ws.cell(row=r, column=8, value=f"={self.gl_sum('Budget', 'ytd', q)}")
                ws.cell(row=r, column=11, value=f"={self.gl_sum('Budget', 'fy', q)}")
                ws.cell(row=r, column=12, value=f"=Forecast!K{self.fc_rows[name]}").font = F_LINK
            else:
                if kind == "total":
                    first, last = rows[item[3][0]], rows[item[3][-1]]
                    f = lambda col: f"=SUM({col}{first}:{col}{last})"
                else:
                    a, b = rows[item[3][0]], rows[item[3][1]]
                    f = lambda col: f"={col}{a}-{col}{b}"
                for c in (3, 4, 7, 8, 11, 12):
                    ws.cell(row=r, column=c, value=f(get_column_letter(c)))
                for c in range(1, 15):
                    ws.cell(row=r, column=c).fill = FILL_TOTAL
                    ws.cell(row=r, column=c).border = TOP
            ws.cell(row=r, column=5, value=self.var(f"$B{r}", f"C{r}", f"D{r}"))
            ws.cell(row=r, column=6, value=self.pct(f"E{r}", f"D{r}"))
            ws.cell(row=r, column=9, value=self.var(f"$B{r}", f"G{r}", f"H{r}"))
            ws.cell(row=r, column=10, value=self.pct(f"I{r}", f"H{r}"))
            ws.cell(row=r, column=13, value=self.var(f"$B{r}", f"L{r}", f"K{r}"))
            ws.cell(row=r, column=14, value=self.flag(f"E{r}", f"F{r}"))
            for c in range(3, 15):
                cell = ws.cell(row=r, column=c)
                cell.number_format = PCT if c in (6, 10) else USD
                if c != 12 or kind != "line":
                    cell.font = F_BOLD if kind != "line" else F_BODY
                if c == 14:
                    cell.alignment = Alignment(horizontal="center")
            rows[name] = r
            r += 1
        self.pnl_rows = rows
        self.flag_rules(ws, f"N6:N{r}")
        ws.freeze_panes = "C6"
        self.widths(ws, [30, 9, 13, 13, 13, 10, 14, 14, 13, 10, 14, 14, 13, 12])
        self.setup_print(ws)

    def dept_tab(self):
        ws = self.wb.create_sheet("Dept Opex", 2)
        self.title(ws, "Operating expenses by department", None)
        ws["A2"] = '="Reporting month: "&Control!D4&"   |   Fav/(Unfav): positive = under budget"'
        ws["A2"].font = F_SUB
        self.head(ws, 4, ["Department / account", "Actual", "Budget", "Fav/(Unfav)", "%", "YTD actual", "YTD budget",
                          "Fav/(Unfav)", "%", "Flag (month)"])
        r = 5
        totals = []
        opex = self.gl_df[self.gl_df["section"] == "Opex"]
        for dept in DEPTS:
            ws.cell(row=r, column=1, value=dept).font = F_BOLD
            for c in range(1, 11):
                ws.cell(row=r, column=c).fill = FILL_SECTION
            ws.cell(row=r, column=11, value=dept)   # helper for the SUMIFS department criterion
            ws.cell(row=r, column=11).font = Font(name=FONT, size=8, color="FFFFFF")
            dept_row = r
            r += 1
            first = r
            for line in [l for l in OPEX if l in set(opex.loc[opex["department"] == dept, "line"])]:
                ws.cell(row=r, column=1, value=line).font = F_BODY
                d = f"$K${dept_row}"
                ws.cell(row=r, column=2, value=f"={self.gl_sum('Actual', 'm', f'$A{r}', d)}")
                ws.cell(row=r, column=3, value=f"={self.gl_sum('Budget', 'm', f'$A{r}', d)}")
                ws.cell(row=r, column=6, value=f"={self.gl_sum('Actual', 'ytd', f'$A{r}', d)}")
                ws.cell(row=r, column=7, value=f"={self.gl_sum('Budget', 'ytd', f'$A{r}', d)}")
                self._dept_vars(ws, r, F_BODY)
                r += 1
            ws.cell(row=r, column=1, value=f"Total {dept}").font = F_BOLD
            for c in (2, 3, 6, 7):
                col = get_column_letter(c)
                ws.cell(row=r, column=c, value=f"=SUM({col}{first}:{col}{r - 1})")
            self._dept_vars(ws, r, F_BOLD, total=True)
            totals.append(r)
            r += 2
        ws.cell(row=r, column=1, value="Total operating expenses").font = F_BOLD
        for c in (2, 3, 6, 7):
            col = get_column_letter(c)
            ws.cell(row=r, column=c, value="=" + "+".join(f"{col}{t}" for t in totals))
        self._dept_vars(ws, r, F_BOLD, total=True)
        self.dept_total_row = r
        self.flag_rules(ws, f"J5:J{r}")
        ws.column_dimensions["K"].hidden = True
        ws.freeze_panes = "B5"
        self.widths(ws, [30, 13, 13, 13, 10, 14, 14, 13, 10, 12])
        self.setup_print(ws)

    def _dept_vars(self, ws, r, font, total=False):
        ws.cell(row=r, column=4, value=f"=C{r}-B{r}")
        ws.cell(row=r, column=5, value=self.pct(f"D{r}", f"C{r}"))
        ws.cell(row=r, column=8, value=f"=G{r}-F{r}")
        ws.cell(row=r, column=9, value=self.pct(f"H{r}", f"G{r}"))
        ws.cell(row=r, column=10, value=self.flag(f"D{r}", f"E{r}"))
        for c in range(2, 11):
            cell = ws.cell(row=r, column=c)
            cell.number_format = PCT if c in (5, 9) else USD
            cell.font = font
            if total:
                cell.fill, cell.border = FILL_TOTAL, TOP
        if total:
            ws.cell(row=r, column=1).fill = FILL_TOTAL
            ws.cell(row=r, column=1).border = TOP
        ws.cell(row=r, column=10).alignment = Alignment(horizontal="center")

    def bridge_tab(self):
        ws = self.wb.create_sheet("Revenue Bridge", 3)
        self.title(ws, "Product revenue bridge: volume, mix and rate",
                   "Volume = change in total TPV at budget mix and rates. Mix = TPV shifted between products at each "
                   "product's budget rate. Rate = actual TPV x (actual - budget take rate).")
        cols = ["Product", "Budget TPV", "Actual TPV", "Budget revenue", "Actual revenue", "Budget rate", "Actual rate",
                "Budget mix", "Volume", "Mix", "Rate", "Total variance", "Check"]
        products = list(PRODUCTS)
        self.bridge_rows = {}
        r = 4
        for label, period in (("Reporting month", "m"), ("Year to date", "ytd")):
            ws.cell(row=r, column=1, value=f"{label}").font = F_BOLD
            ws.cell(row=r, column=2, value='=Control!D4' if period == "m" else '="Jan - "&Control!D4').font = F_LINK
            r += 1
            self.head(ws, r, cols)
            r += 1
            first = r
            last = r + len(products) - 1
            tot = last + 1
            for p in products:
                ws.cell(row=r, column=1, value=p).font = F_BODY
                dcrit = lambda col: f'SUMIFS({self.drv(col)},{self.drv("product")},$A{r},{",".join(self._period(self.drv("month_num"), period))})'
                ws.cell(row=r, column=2, value=f"={dcrit('budget_tpv')}")
                ws.cell(row=r, column=3, value=f"={dcrit('actual_tpv')}")
                ws.cell(row=r, column=4, value=f"={self.gl_sum('Budget', period, f'$A{r}')}")
                ws.cell(row=r, column=5, value=f"={self.gl_sum('Actual', period, f'$A{r}')}")
                ws.cell(row=r, column=6, value=f"=IFERROR(D{r}/B{r},0)")
                ws.cell(row=r, column=7, value=f"=IFERROR(E{r}/C{r},0)")
                ws.cell(row=r, column=8, value=f"=IFERROR(B{r}/$B${tot},0)")
                ws.cell(row=r, column=9, value=f"=($C${tot}-$B${tot})*H{r}*F{r}")
                ws.cell(row=r, column=10, value=f"=(C{r}-$C${tot}*H{r})*F{r}")
                ws.cell(row=r, column=11, value=f"=C{r}*(G{r}-F{r})")
                ws.cell(row=r, column=12, value=f"=E{r}-D{r}")
                ws.cell(row=r, column=13, value=f"=ROUND(I{r}+J{r}+K{r}-L{r},2)")
                r += 1
            ws.cell(row=r, column=1, value="Total products").font = F_BOLD
            for c in (2, 3, 4, 5, 8, 9, 10, 11, 12, 13):
                col = get_column_letter(c)
                ws.cell(row=r, column=c, value=f"=SUM({col}{first}:{col}{last})")
            ws.cell(row=r, column=6, value=f"=IFERROR(D{r}/B{r},0)")
            ws.cell(row=r, column=7, value=f"=IFERROR(E{r}/C{r},0)")
            for rr in range(first, r + 1):
                for c in range(2, 14):
                    cell = ws.cell(row=rr, column=c)
                    cell.number_format = RATE if c in (6, 7) else ("0.0%" if c == 8 else USD)
                    cell.font = F_BOLD if rr == r else F_BODY
                    if rr == r:
                        cell.fill, cell.border = FILL_TOTAL, TOP
            self.bridge_rows[period] = (first, r)
            r += 3

        # waterfall (YTD): invisible base + up + down series
        first, tot = self.bridge_rows["ytd"]
        ws.cell(row=r, column=1, value="YTD revenue variance waterfall ($)").font = F_BOLD
        r += 1
        self.head(ws, r, ["Step", "Value", "Running total", "Base", "Increase", "Decrease"])
        steps = [("Volume", f"=I{tot}"), ("Mix", f"=J{tot}"), ("Rate", f"=K{tot}"), ("Total variance", f"=L{tot}")]
        start = r + 1
        for i, (name, val) in enumerate(steps):
            rr = start + i
            ws.cell(row=rr, column=1, value=name).font = F_BODY
            ws.cell(row=rr, column=2, value=val)
            if name == "Total variance":
                ws.cell(row=rr, column=3, value=f"=B{rr}")
                ws.cell(row=rr, column=4, value=f"=MIN(0,B{rr})")
                ws.cell(row=rr, column=5, value=f"=MAX(0,B{rr})")
                ws.cell(row=rr, column=6, value=f"=MAX(0,-B{rr})")
            else:
                prev = f"C{rr - 1}" if i else "0"
                ws.cell(row=rr, column=3, value=f"={prev}+B{rr}")
                ws.cell(row=rr, column=4, value=f"=MIN({prev},C{rr})")
                ws.cell(row=rr, column=5, value=f"=MAX(0,B{rr})")
                ws.cell(row=rr, column=6, value=f"=MAX(0,-B{rr})")
            for c in range(2, 7):
                ws.cell(row=rr, column=c).number_format = USD
                ws.cell(row=rr, column=c).font = F_BODY
        end = start + len(steps) - 1
        chart = BarChart()
        chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
        chart.title = "YTD revenue variance bridge"
        chart.y_axis.title = "$"
        chart.y_axis.numFmt = '$#,##0,"K"'
        data = Reference(ws, min_col=4, max_col=6, min_row=start - 1, max_row=end)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=start, max_row=end))
        chart.series[0].graphicalProperties.noFill = True
        chart.series[0].graphicalProperties.line.noFill = True
        chart.series[1].graphicalProperties.solidFill = "2A78D6"
        chart.series[2].graphicalProperties.solidFill = "D03B3B"
        from openpyxl.chart.legend import LegendEntry
        chart.legend.position = "b"
        chart.legend.legendEntry = [LegendEntry(idx=0, delete=True)]
        chart.height, chart.width = 8, 16
        ws.add_chart(chart, f"H{r - 1}")
        ws.freeze_panes = "B4"
        self.widths(ws, [30, 16, 16, 15, 15, 11, 11, 10, 13, 13, 13, 14, 8])
        self.setup_print(ws)

    def headcount_tab(self):
        ws = self.wb.create_sheet("Headcount", 4)
        self.title(ws, "Headcount and salary variance",
                   "Volume effect = (budget - actual head-months YTD) x budget cost per head. Rate effect = the rest "
                   "(pay, benefits and hiring mix).")
        ws["A3"] = '="Headcount at "&Control!D4&"; salaries Jan to "&Control!D4'
        ws["A3"].font = F_SUB
        self.head(ws, 5, ["Department", "Budget HC", "Actual HC", "HC gap", "Budget cost / head", "Budget head-months",
                          "Actual head-months", "YTD budget salaries", "YTD actual salaries", "Fav/(Unfav)",
                          "Volume (headcount)", "Rate (cost / head)"])
        r = 6
        for dept in DEPTS:
            ws.cell(row=r, column=1, value=dept).font = F_BODY
            hm = lambda col, per: f'SUMIFS({self.hcr(col)},{self.hcr("department")},$A{r},{",".join(self._period(self.hcr("month_num"), per))})'
            ws.cell(row=r, column=2, value=f"={hm('budget_hc', 'm')}")
            ws.cell(row=r, column=3, value=f"={hm('actual_hc', 'm')}")
            ws.cell(row=r, column=4, value=f"=C{r}-B{r}")
            ws.cell(row=r, column=5, value=f'=AVERAGEIFS({self.hcr("cph")},{self.hcr("department")},$A{r})')
            ws.cell(row=r, column=6, value=f"={hm('budget_hc', 'ytd')}")
            ws.cell(row=r, column=7, value=f"={hm('actual_hc', 'ytd')}")
            ws.cell(row=r, column=8, value=f"={self.gl_sum('Budget', 'ytd', chr(34) + 'Salaries & benefits' + chr(34), f'$A{r}')}")
            ws.cell(row=r, column=9, value=f"={self.gl_sum('Actual', 'ytd', chr(34) + 'Salaries & benefits' + chr(34), f'$A{r}')}")
            ws.cell(row=r, column=10, value=f"=H{r}-I{r}")
            ws.cell(row=r, column=11, value=f"=(F{r}-G{r})*E{r}")
            ws.cell(row=r, column=12, value=f"=J{r}-K{r}")
            r += 1
        ws.cell(row=r, column=1, value="Total").font = F_BOLD
        for c in (2, 3, 4, 6, 7, 8, 9, 10, 11, 12):
            col = get_column_letter(c)
            ws.cell(row=r, column=c, value=f"=SUM({col}6:{col}{r - 1})")
        for rr in range(6, r + 1):
            for c in range(2, 13):
                cell = ws.cell(row=rr, column=c)
                cell.number_format = NUM if c in (2, 3, 4, 6, 7) else USD
                cell.font = F_BOLD if rr == r else F_BODY
                if rr == r:
                    cell.fill, cell.border = FILL_TOTAL, TOP
        self.hc_total_row = r
        self.widths(ws, [22, 11, 11, 9, 14, 14, 14, 16, 16, 14, 15, 15])
        self.setup_print(ws)

    def commentary_tab(self):
        ws = self.wb.create_sheet("Commentary", 6)
        self.title(ws, "September 2026 close: commentary",
                   "Written at the September close. These notes do not change if the reporting month is changed.")
        r = 4
        for heading, bullets in self.commentary:
            ws.cell(row=r, column=1, value=heading).font = F_BOLD
            r += 1
            for b in bullets:
                c = ws.cell(row=r, column=1, value="•  " + b)
                c.font = F_BODY
                c.alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[r].height = 15 * max(1, len(b) // 120 + 1)
                r += 1
            r += 1
        self.widths(ws, [130])

    def build(self, path):
        self.control()
        self.data_tabs()
        self.forecast_tab()
        self.pnl_tab()
        self.dept_tab()
        self.bridge_tab()
        self.headcount_tab()
        self.commentary_tab()
        order = ["Control", "P&L Summary", "Dept Opex", "Revenue Bridge", "Headcount", "Forecast", "Commentary",
                 "GL_Data", "Drivers", "HC_Data"]
        self.wb._sheets = [self.wb[n] for n in order]
        self.wb.active = 1
        self.wb.save(path)
        return path
